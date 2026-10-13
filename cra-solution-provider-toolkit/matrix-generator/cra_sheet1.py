import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from cra_styles import *

# -------------------------------------------------------------
# SHEET 1: CRA REQUIREMENTS MATRIX
# -------------------------------------------------------------
def build_requirements_matrix(wb, reqs):
    print("Building Sheet 1: CRA Requirements Matrix...")
    ws = wb.create_sheet(title="CRA Requirements Matrix")
    ws.views.sheetView[0].showGridLines = True
    
    # Title Block
    ws.merge_cells("A1:AE1")
    ws["A1"] = "EU CYBER RESILIENCE ACT (REGULATION (EU) 2024/2847) — COMPLIANCE REQUIREMENTS MATRIX"
    ws["A1"].font = font_title
    ws["A1"].alignment = Alignment(vertical="center", horizontal="left")
    ws.row_dimensions[1].height = 25
    
    ws.merge_cells("A2:AE2")
    ws["A2"] = "Coverage Definitions: Full (Direct principal technical evidence within scope) | Partial (Supporting evidence; requires complementary controls) | None (No material direct contribution) | Baseline: Standard Product with Digital Elements placed on EU Market"
    ws["A2"].font = font_subtitle
    ws["A2"].alignment = Alignment(vertical="center", horizontal="left")
    ws.row_dimensions[2].height = 18

    headers = [
        "Requirement ID", "CRA Legal Reference", "Requirement Title", "Requirement Description & Legal Intent",
        "Requirement Type", "Applicability", "Responsible Stakeholder", "Presales / Planning",
        "Development", "Build / Release", "Deployment", "Maintenance / RUN",
        "Implementation Guidance", "Recommended Technical & Organisational Controls",
        "Required Evidence & Compliance Deliverables", "Verification or Acceptance Criteria",
        "SonarQube Coverage", "Nexus IQ CLM Coverage", "Penetration Testing Coverage", "Nessus / Cyberwatch Coverage",
        "DPIA Coverage", "Threat Analysis / Risk Analysis Coverage", "Dependency-Track Coverage", "Black Duck Coverage",
        "Coverity Coverage", "OWASP ZAP / Burp Coverage", "Coverage Limitations & Assumptions",
        "Additional Controls or Activities Required", "Remediation Priority", "Official Source URL",
        "Notes & Product-Specific Applicability Considerations"
    ]
    
    ws.row_dimensions[3].height = 28
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header_navy
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border_thin

    row_start = 4
    for r_idx, r in enumerate(reqs, start=row_start):
        ws.row_dimensions[r_idx].height = 65
        is_zebra = (r_idx % 2 == 0)
        base_fill = fill_zebra if is_zebra else fill_white
        
        row_values = [
            r["id"], r["ref"], r["title"], r["description"],
            r["req_type"], r["applicability"], r["stakeholder"],
            r["presales"], r["dev"], r["build"], r["deploy"], r["run"],
            r["guidance"], r["controls"], r["evidence"], r["criteria"],
            r["sonar"], r["nexus"], r["pentest"], r["nessus"],
            r["dpia"], r["tara"], r["deptrack"], r["blackduck"],
            r["coverity"], r["dast"], r["limitations"], r["additional_controls"],
            r["priority"], r["url"], r["notes"]
        ]
        
        for col_idx, val in enumerate(row_values, 1):
            cell = ws.cell(row=r_idx, column=col_idx, value=val)
            cell.font = font_cell
            cell.border = border_thin
            cell.fill = base_fill
            
            # Alignments & Styling based on column
            if col_idx == 1: # Req ID
                cell.font = font_cell_bold
                cell.alignment = Alignment(horizontal="center", vertical="top")
            elif col_idx in [2, 5, 7]: # Ref, Type, Stakeholder
                cell.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
            elif col_idx == 6: # Applicability
                cell.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
                if val.startswith("Mandatory"):
                    cell.fill = fill_mand
                    cell.font = font_mand
                elif val.startswith("Conditional"):
                    cell.fill = fill_cond
                    cell.font = font_cond
            elif col_idx in [8, 9, 10, 11, 12]: # Lifecycle phases
                cell.alignment = Alignment(horizontal="center", vertical="top")
                if val == "Yes":
                    cell.fill = fill_yes
                    cell.font = font_yes
                elif val == "Supporting":
                    cell.fill = fill_supp
                    cell.font = font_supp
                elif val == "No":
                    cell.fill = fill_no
                    cell.font = font_no
            elif 17 <= col_idx <= 26: # Tool coverages
                cell.alignment = Alignment(horizontal="center", vertical="top")
                if val == "Full":
                    cell.fill = fill_full
                    cell.font = font_full
                elif val == "Partial":
                    cell.fill = fill_partial
                    cell.font = font_partial
                elif val == "None":
                    cell.fill = fill_none
                    cell.font = font_none
            elif col_idx == 29: # Priority
                cell.alignment = Alignment(horizontal="center", vertical="top")
                if val == "Critical":
                    cell.fill = fill_crit
                    cell.font = font_crit
                elif val == "High":
                    cell.fill = fill_high
                    cell.font = font_high
                elif val == "Medium":
                    cell.fill = fill_med
                    cell.font = font_med
                elif val == "Low":
                    cell.fill = fill_low
                    cell.font = font_low
            elif col_idx == 30: # URL
                cell.hyperlink = val
                cell.font = font_link
                cell.alignment = Alignment(horizontal="center", vertical="top")
            else: # Text descriptions, guidance, evidence, notes
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

    # Freeze panes: freeze rows 1-3, and columns A-C
    ws.freeze_panes = "D4"
    
    # Auto-filter
    ws.auto_filter.ref = f"A3:AE{len(reqs) + 3}"
    
    # Data Validations
    dv_coverage = DataValidation(type="list", formula1='"Full,Partial,None"', allow_blank=True)
    ws.add_data_validation(dv_coverage)
    dv_coverage.add(f"Q4:Z{len(reqs) + 3}")
    
    dv_priority = DataValidation(type="list", formula1='"Critical,High,Medium,Low"', allow_blank=True)
    ws.add_data_validation(dv_priority)
    dv_priority.add(f"AC4:AC{len(reqs) + 3}")
    
    dv_lifecycle = DataValidation(type="list", formula1='"Yes,Supporting,No"', allow_blank=True)
    ws.add_data_validation(dv_lifecycle)
    dv_lifecycle.add(f"H4:L{len(reqs) + 3}")
    
    # Custom column widths
    custom_widths = {
        "A": 16, "B": 24, "C": 30, "D": 45, "E": 20, "F": 22, "G": 24,
        "H": 14, "I": 14, "J": 14, "K": 14, "L": 15, "M": 45, "N": 45,
        "O": 45, "P": 45, "Q": 14, "R": 14, "S": 15, "T": 15, "U": 14,
        "V": 16, "W": 15, "X": 14, "Y": 14, "Z": 14, "AA": 42, "AB": 42,
        "AC": 15, "AD": 25, "AE": 40
    }
    auto_fit_columns(ws, custom_widths=custom_widths)
    print("Sheet 1 built successfully.")
