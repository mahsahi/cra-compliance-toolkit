import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from cra_styles import *

# -------------------------------------------------------------
# SHEET 5: SOURCES AND ASSUMPTIONS
# -------------------------------------------------------------
def build_sources_and_assumptions(wb, sources_data):
    print("Building Sheet 5: Sources and Assumptions...")
    ws = wb.create_sheet(title="Sources and Assumptions")
    ws.views.sheetView[0].showGridLines = True
    
    # Title Block
    ws.merge_cells("A1:G1")
    ws["A1"] = "AUTHORITATIVE REGULATORY SOURCES, STANDARDS REGISTER & COMPLIANCE ASSUMPTIONS"
    ws["A1"].font = font_title
    ws["A1"].alignment = Alignment(vertical="center", horizontal="left")
    ws.row_dimensions[1].height = 25
    
    ws.merge_cells("A2:G2")
    ws["A2"] = f"Regulatory Baseline: Regulation (EU) 2024/2847 | Assessment Date: {sources_data['review_metadata']['date_of_review']} | Status: {sources_data['review_metadata']['regulatory_status']}"
    ws["A2"].font = font_subtitle
    ws["A2"].alignment = Alignment(vertical="center", horizontal="left")
    ws.row_dimensions[2].height = 18

    curr_row = 4

    # 1. Official Regulations
    ws.merge_cells(f"A{curr_row}:G{curr_row}")
    ws[f"A{curr_row}"] = "1. PRIMARY REGULATORY FRAMEWORK & BINDING UNION ACTS"
    ws[f"A{curr_row}"].font = font_section
    ws[f"A{curr_row}"].fill = fill_dark_navy
    ws[f"A{curr_row}"].alignment = Alignment(vertical="center", horizontal="left")
    ws.row_dimensions[curr_row].height = 22
    curr_row += 1

    headers_reg = ["Legal Instrument", "Official Title", "OJ Citation", "Entry into Force", "Application Date", "Legal Nature", "Official Source URL"]
    ws.row_dimensions[curr_row].height = 25
    for c_idx, h in enumerate(headers_reg, 1):
        cell = ws.cell(row=curr_row, column=c_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header_navy
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border_thin
    curr_row += 1

    for r in sources_data["regulations"]:
        ws.row_dimensions[curr_row].height = 35
        vals = [r["instrument"], r["official_title"], r["citation"], r["entry_into_force"], r["applicability_date"], r["legal_nature"], r["url"]]
        for c_idx, val in enumerate(vals, 1):
            cell = ws.cell(row=curr_row, column=c_idx, value=val)
            cell.font = font_cell
            cell.border = border_thin
            if c_idx == 1:
                cell.font = font_cell_bold
                cell.alignment = Alignment(horizontal="left", vertical="top")
            elif c_idx == 7:
                cell.hyperlink = val
                cell.font = font_link
                cell.alignment = Alignment(horizontal="center", vertical="top")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        curr_row += 1

    curr_row += 1
    # 2. Authoritative Guidance
    ws.merge_cells(f"A{curr_row}:G{curr_row}")
    ws[f"A{curr_row}"] = "2. AUTHORITATIVE EUROPEAN COMMISSION GUIDANCE & ADMINISTRATIVE NOTICES"
    ws[f"A{curr_row}"].font = font_section
    ws[f"A{curr_row}"].fill = fill_dark_navy
    ws[f"A{curr_row}"].alignment = Alignment(vertical="center", horizontal="left")
    ws.row_dimensions[curr_row].height = 22
    curr_row += 1

    headers_guid = ["Document Title", "Issuing Body", "Publication Date", "Legal Status", "Key Subject Matter & Coverage", "Official Source URL", "Notes"]
    ws.row_dimensions[curr_row].height = 25
    for c_idx, h in enumerate(headers_guid, 1):
        cell = ws.cell(row=curr_row, column=c_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header_navy
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border_thin
    curr_row += 1

    for g in sources_data["guidance"]:
        ws.row_dimensions[curr_row].height = 35
        vals = [g["document"], g["issuing_body"], g["publication_date"], g["status"], g["key_topics"], g["url"], "Authoritative EC interpretation"]
        for c_idx, val in enumerate(vals, 1):
            cell = ws.cell(row=curr_row, column=c_idx, value=val)
            cell.font = font_cell
            cell.border = border_thin
            if c_idx == 1:
                cell.font = font_cell_bold
                cell.alignment = Alignment(horizontal="left", vertical="top")
            elif c_idx == 6:
                cell.hyperlink = val
                cell.font = font_link
                cell.alignment = Alignment(horizontal="center", vertical="top")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        curr_row += 1

    curr_row += 1
    # 3. Standards Register
    ws.merge_cells(f"A{curr_row}:G{curr_row}")
    ws[f"A{curr_row}"] = "3. HARMONISED STANDARDS & TECHNICAL BENCHMARKS REGISTER"
    ws[f"A{curr_row}"].font = font_section
    ws[f"A{curr_row}"].fill = fill_dark_navy
    ws[f"A{curr_row}"].alignment = Alignment(vertical="center", horizontal="left")
    ws.row_dimensions[curr_row].height = 22
    curr_row += 1

    headers_std = ["Standard Reference", "Standard Title", "Current Publication Status", "Verified Harmonised Status under CRA", "Relevance to CRA Requirements", "State-of-the-Art Audit Notes", "Candidate for Harmonisation"]
    ws.row_dimensions[curr_row].height = 25
    for c_idx, h in enumerate(headers_std, 1):
        cell = ws.cell(row=curr_row, column=c_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header_slate
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border_thin
    curr_row += 1

    for s in sources_data["standards"]:
        ws.row_dimensions[curr_row].height = 38
        vals = [s["standard_ref"], s["title"], s["current_status"], s["verified_harmonised_status"], s["relevance"], s["notes"], "Yes (Future Presumption of Conformity)"]
        for c_idx, val in enumerate(vals, 1):
            cell = ws.cell(row=curr_row, column=c_idx, value=val)
            cell.font = font_cell
            cell.border = border_thin
            if c_idx == 1:
                cell.font = font_cell_bold
                cell.alignment = Alignment(horizontal="left", vertical="top")
            elif c_idx == 4: # Harmonised status
                cell.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
                if "NOT HARMONISED" in val:
                    cell.fill = fill_crit
                    cell.font = font_crit
            else:
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        curr_row += 1

    curr_row += 1
    # 4. Product Assumptions
    ws.merge_cells(f"A{curr_row}:G{curr_row}")
    ws[f"A{curr_row}"] = "4. PRODUCT SCOPE & BASELINE ARCHITECTURE ASSUMPTIONS"
    ws[f"A{curr_row}"].font = font_section
    ws[f"A{curr_row}"].fill = fill_dark_navy
    ws[f"A{curr_row}"].alignment = Alignment(vertical="center", horizontal="left")
    ws.row_dimensions[curr_row].height = 22
    curr_row += 1

    headers_assump = ["Architectural / Legal Area", "Assumption Description", "CRA Regulatory Basis", "Operational & Compliance Impact", "Validation Requirement", "Risk Level if Invalidated", "Owner"]
    ws.row_dimensions[curr_row].height = 25
    for c_idx, h in enumerate(headers_assump, 1):
        cell = ws.cell(row=curr_row, column=c_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header_navy
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border_thin
    curr_row += 1

    for a in sources_data["assumptions"]:
        ws.row_dimensions[curr_row].height = 38
        vals = [a["area"], a["assumption_text"], a["regulatory_basis"], a["impact_on_findings"], "Verify prior to project kickoff", "High", "Legal / Lead Architect"]
        for c_idx, val in enumerate(vals, 1):
            cell = ws.cell(row=curr_row, column=c_idx, value=val)
            cell.font = font_cell
            cell.border = border_thin
            if c_idx == 1:
                cell.font = font_cell_bold
                cell.alignment = Alignment(horizontal="left", vertical="top")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        curr_row += 1

    curr_row += 1
    # 5. Known Uncertainties
    ws.merge_cells(f"A{curr_row}:G{curr_row}")
    ws[f"A{curr_row}"] = "5. KNOWN REGULATORY UNCERTAINTIES & ITEMS REQUIRING SPECIALIZED LEGAL/NOTIFIED BODY REVIEW"
    ws[f"A{curr_row}"].font = font_section
    ws[f"A{curr_row}"].fill = fill_dark_navy
    ws[f"A{curr_row}"].alignment = Alignment(vertical="center", horizontal="left")
    ws.row_dimensions[curr_row].height = 22
    curr_row += 1

    headers_unc = ["Uncertainty / Ambiguity Area", "Detailed Description & Legal Complexity", "Recommended Precautionary Measure", "Reviewing Stakeholder", "Monitoring Mechanism", "Escalation Trigger", "Impact on CE Mark"]
    ws.row_dimensions[curr_row].height = 25
    for c_idx, h in enumerate(headers_unc, 1):
        cell = ws.cell(row=curr_row, column=c_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header_slate
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border_thin
    curr_row += 1

    for u in sources_data["uncertainties"]:
        ws.row_dimensions[curr_row].height = 38
        vals = [u["issue"], u["description"], u["recommended_action"], "General Counsel & CISO", "Quarterly Regulatory Watch", "Formal EC Implementing Act publication", "Critical"]
        for c_idx, val in enumerate(vals, 1):
            cell = ws.cell(row=curr_row, column=c_idx, value=val)
            cell.font = font_cell
            cell.border = border_thin
            if c_idx == 1:
                cell.font = font_cell_bold
                cell.alignment = Alignment(horizontal="left", vertical="top")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        curr_row += 1

    # Freeze panes
    ws.freeze_panes = "A4"

    # Widths
    custom_widths = {
        "A": 28, "B": 38, "C": 22, "D": 22, "E": 30, "F": 28, "G": 30
    }
    auto_fit_columns(ws, custom_widths=custom_widths)
    print("Sheet 5 built successfully.")
