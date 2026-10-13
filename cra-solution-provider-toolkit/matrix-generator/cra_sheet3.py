import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from cra_styles import *

# -------------------------------------------------------------
# SHEET 3: LIFECYCLE CONTROL MAPPING
# -------------------------------------------------------------
def build_lifecycle_mapping(wb, lifecycle_mappings):
    print("Building Sheet 3: Lifecycle Control Mapping...")
    ws = wb.create_sheet(title="Lifecycle Control Mapping")
    ws.views.sheetView[0].showGridLines = True
    
    # Title Block
    ws.merge_cells("A1:H1")
    ws["A1"] = "CRA LIFECYCLE CONTROL MAPPING & OPERATIONAL CONTROL GATE ARCHITECTURE"
    ws["A1"].font = font_title
    ws["A1"].alignment = Alignment(vertical="center", horizontal="left")
    ws.row_dimensions[1].height = 25
    
    ws.merge_cells("A2:H2")
    ws["A2"] = "Operational Security Actions, Supporting Tools, Required Deliverables, and Approval Gates across 5 Lifecycle Phases"
    ws["A2"].font = font_subtitle
    ws["A2"].alignment = Alignment(vertical="center", horizontal="left")
    ws.row_dimensions[2].height = 18

    headers = [
        "CRA Legal Reference", "Lifecycle Phase", "Required Action", "Responsible Role",
        "Tool or Process Supporting Action", "Required Input and Output",
        "Evidence to Retain", "Suggested Control Gate or Approval Criterion"
    ]
    
    ws.row_dimensions[3].height = 26
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header_navy
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border_thin

    phase_fills = {
        "Presales and Product Planning": PatternFill(start_color="E0F2FE", end_color="E0F2FE", fill_type="solid"),
        "Development": PatternFill(start_color="EDE9FE", end_color="EDE9FE", fill_type="solid"),
        "Build and Release": PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid"),
        "Deployment and Delivery": PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid"),
        "Maintenance and RUN": PatternFill(start_color="FCE7F3", end_color="FCE7F3", fill_type="solid")
    }
    phase_fonts = {
        "Presales and Product Planning": Font(name=FONT_FAMILY, size=9, bold=True, color="0369A1"),
        "Development": Font(name=FONT_FAMILY, size=9, bold=True, color="5B21B6"),
        "Build and Release": Font(name=FONT_FAMILY, size=9, bold=True, color="92400E"),
        "Deployment and Delivery": Font(name=FONT_FAMILY, size=9, bold=True, color="166534"),
        "Maintenance and RUN": Font(name=FONT_FAMILY, size=9, bold=True, color="9D174D")
    }

    for idx, item in enumerate(lifecycle_mappings, start=4):
        ws.row_dimensions[idx].height = 45
        is_zebra = (idx % 2 == 0)
        base_fill = fill_zebra if is_zebra else fill_white
        
        row_vals = [
            item["cra_ref"], item["phase"], item["action"], item["role"],
            item["tool_process"], item["input_output"], item["evidence"], item["gate"]
        ]
        
        for c_idx, val in enumerate(row_vals, 1):
            cell = ws.cell(row=idx, column=c_idx, value=val)
            cell.font = font_cell
            cell.border = border_thin
            cell.fill = base_fill
            
            if c_idx == 1:
                cell.alignment = Alignment(horizontal="center", vertical="top")
                cell.font = font_cell_bold
            elif c_idx == 2:
                cell.alignment = Alignment(horizontal="center", vertical="top")
                p_name = item["phase"]
                if p_name in phase_fills:
                    cell.fill = phase_fills[p_name]
                    cell.font = phase_fonts[p_name]
            elif c_idx in [3, 4]:
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
                if c_idx == 3:
                    cell.font = font_cell_bold
            else:
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

    # Freeze panes
    ws.freeze_panes = "A4"
    
    # Auto-filter
    ws.auto_filter.ref = f"A3:H{len(lifecycle_mappings) + 3}"
    
    # Widths
    custom_widths = {
        "A": 22, "B": 24, "C": 35, "D": 25, "E": 30, "F": 45, "G": 42, "H": 45
    }
    auto_fit_columns(ws, custom_widths=custom_widths)
    print("Sheet 3 built successfully.")
