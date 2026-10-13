import json
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Load datasets
with open('all_reqs.json') as f:
    reqs = json.load(f)
with open('tool_profiles.json') as f:
    tool_profiles = json.load(f)
with open('lifecycle_mappings.json') as f:
    lifecycle_mappings = json.load(f)
with open('gaps_roadmap.json') as f:
    gaps_roadmap = json.load(f)
with open('sources_data.json') as f:
    sources_data = json.load(f)

# Import builder functions
from cra_sheet1 import build_requirements_matrix
from cra_sheet2 import build_tool_summary
from cra_sheet3 import build_lifecycle_mapping
from cra_sheet4 import build_gaps_roadmap
from cra_sheet5 import build_sources_and_assumptions
from cra_dashboard import build_executive_dashboard

wb = openpyxl.Workbook()
wb.remove(wb.active) # remove default sheet

# Build sheets in exact requested order
# 1. Executive Dashboard (Executive overview & KPI cards)
build_executive_dashboard(wb, reqs, tool_profiles, lifecycle_mappings, gaps_roadmap)

# 2. Sheet 1 — CRA Requirements Matrix
build_requirements_matrix(wb, reqs)

# 3. Sheet 2 — Tool Coverage Summary
build_tool_summary(wb, tool_profiles)

# 4. Sheet 3 — Lifecycle Control Mapping
build_lifecycle_mapping(wb, lifecycle_mappings)

# 5. Sheet 4 — Compliance Gaps and Remediation Roadmap
build_gaps_roadmap(wb, gaps_roadmap)

# 6. Sheet 5 — Sources and Assumptions
build_sources_and_assumptions(wb, sources_data)

target_filename = "CRA_Compliance_Requirements_and_Security_Tool_Coverage_Matrix.xlsx"
print(f"Saving completed workbook to {target_filename}...")
wb.save(target_filename)
file_size = os.path.getsize(target_filename)
print(f"SUCCESS: Workbook saved successfully! File size: {file_size:,} bytes.")
