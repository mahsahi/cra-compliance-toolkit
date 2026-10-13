import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from cra_styles import *

# -------------------------------------------------------------
# SHEET 4: COMPLIANCE GAPS AND REMEDIATION ROADMAP
# -------------------------------------------------------------
def build_gaps_roadmap(wb, gaps_roadmap):
    print("Building Sheet 4: Compliance Gaps & Roadmap...")
    ws = wb.create_sheet(title="Compliance Gaps & Roadmap")
    ws.views.sheetView[0].showGridLines = True
    
    # Title Block
    ws.merge_cells("A1:I1")
    ws["A1"] = "CRA COMPLIANCE GAPS & STRATEGIC REMEDIATION ROADMAP"
    ws["A1"].font = font_title
    ws["A1"].alignment = Alignment(vertical="center", horizontal="left")
    ws.row_dimensions[1].height = 25
    
    ws.merge_cells("A2:I2")
    ws["A2"] = "Detailed Analysis of Residual Regulatory Gaps Not Fully Addressed by the Security Toolchain, Requiring Governance, Legal, and Operational Controls"
    ws["A2"].font = font_subtitle
    ws["A2"].alignment = Alignment(vertical="center", horizontal="left")
    ws.row_dimensions[2].height = 18

    headers = [
        "Gap ID", "Mapped CRA Requirement ID", "Identified Compliance Gap",
        "Regulatory Risk & Business Impact", "Recommended Complementary Control / Process",
        "Responsible Owner", "Priority & Rationale", "Suggested Implementation Phase",
        "Required Compliance Evidence & Dependencies"
    ]
    
    ws.row_dimensions[3].height = 26
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header_navy
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border_thin

    for idx, item in enumerate(gaps_roadmap, start=4):
        ws.row_dimensions[idx].height = 55
        is_zebra = (idx % 2 == 0)
        base_fill = fill_zebra if is_zebra else fill_white
        
        row_vals = [
            item["id"], item["req_id"], item["gap"], item["risk"],
            item["control"], item["owner"], item["priority"], item["phase"],
            f"Evidence: {item['evidence']} | Dependencies: {item['dependencies']}"
        ]
        
        for c_idx, val in enumerate(row_vals, 1):
            cell = ws.cell(row=idx, column=c_idx, value=val)
            cell.font = font_cell
            cell.border = border_thin
            cell.fill = base_fill
            
            if c_idx in [1, 2]:
                cell.alignment = Alignment(horizontal="center", vertical="top")
                cell.font = font_cell_bold
            elif c_idx == 6:
                cell.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
            elif c_idx == 7:
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
                if val.startswith("Critical"):
                    cell.fill = fill_crit
                    cell.font = font_crit
                elif val.startswith("High"):
                    cell.fill = fill_high
                    cell.font = font_high
                elif val.startswith("Medium"):
                    cell.fill = fill_med
                    cell.font = font_med
                elif val.startswith("Low"):
                    cell.fill = fill_low
                    cell.font = font_low
            elif c_idx == 8:
                cell.alignment = Alignment(horizontal="center", vertical="top")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

    # Freeze panes
    ws.freeze_panes = "A4"
    
    # Auto-filter
    ws.auto_filter.ref = f"A3:I{len(gaps_roadmap) + 3}"
    
    # Widths
    custom_widths = {
        "A": 16, "B": 22, "C": 38, "D": 38, "E": 40, "F": 25, "G": 35, "H": 24, "I": 45
    }
    auto_fit_columns(ws, custom_widths=custom_widths)
    print("Sheet 4 built successfully.")
