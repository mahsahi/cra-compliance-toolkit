import json
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule

# Load all prepared datasets
print("Loading JSON datasets...")
with open('all_reqs.json') as f:
    reqs = json.load(f)
with open('tool_profiles.json') as f:
    tool_profiles = json.load(f)
with open('lifecycle_mappings.json') as f:
    lifecycle_mappings = json.load(f)
with open('gaps_roadmap.json') as f:
    gaps_roadmap = json.load(f)
with open('sources_data.json') as f:
    sources_data = json.load(f)

print(f"Loaded: {len(reqs)} reqs, {len(tool_profiles)} tools, {len(lifecycle_mappings)} lifecycle actions, {len(gaps_roadmap)} gaps.")

# Initialize Workbook
wb = openpyxl.Workbook()
wb.remove(wb.active) # Remove default sheet

# Common Styling Definitions
FONT_FAMILY = "Segoe UI"

font_title = Font(name=FONT_FAMILY, size=15, bold=True, color="0F2537")
font_subtitle = Font(name=FONT_FAMILY, size=9.5, italic=True, color="4A5568")
font_section = Font(name=FONT_FAMILY, size=11, bold=True, color="FFFFFF")
font_subheading = Font(name=FONT_FAMILY, size=10, bold=True, color="0F2537")
font_header = Font(name=FONT_FAMILY, size=9.5, bold=True, color="FFFFFF")
font_cell = Font(name=FONT_FAMILY, size=9, color="1A202C")
font_cell_bold = Font(name=FONT_FAMILY, size=9, bold=True, color="1A202C")
font_kpi_num = Font(name=FONT_FAMILY, size=18, bold=True, color="0F2537")
font_kpi_label = Font(name=FONT_FAMILY, size=8.5, bold=True, color="4A5568")
font_link = Font(name=FONT_FAMILY, size=9, underline="single", color="1B44BA")

# Fills
fill_dark_navy = PatternFill(start_color="0F2537", end_color="0F2537", fill_type="solid")
fill_header_navy = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
fill_header_slate = PatternFill(start_color="2C5282", end_color="2C5282", fill_type="solid")
fill_section_ice = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
fill_card_bg = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

# Status / Metric Fills & Fonts
fill_full = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid")
font_full = Font(name=FONT_FAMILY, size=9, bold=True, color="155724")

fill_partial = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")
font_partial = Font(name=FONT_FAMILY, size=9, bold=True, color="856404")

fill_none = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")
font_none = Font(name=FONT_FAMILY, size=9, color="721C24")

fill_mand = PatternFill(start_color="E0F2FE", end_color="E0F2FE", fill_type="solid")
font_mand = Font(name=FONT_FAMILY, size=9, color="0369A1")

fill_cond = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
font_cond = Font(name=FONT_FAMILY, size=9, color="B45309")

fill_crit = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
font_crit = Font(name=FONT_FAMILY, size=9, bold=True, color="991B1B")

fill_high = PatternFill(start_color="FFEDD5", end_color="FFEDD5", fill_type="solid")
font_high = Font(name=FONT_FAMILY, size=9, bold=True, color="C2410C")

fill_med = PatternFill(start_color="FEF9C3", end_color="FEF9C3", fill_type="solid")
font_med = Font(name=FONT_FAMILY, size=9, color="854D0E")

fill_low = PatternFill(start_color="F0FDF4", end_color="F0FDF4", fill_type="solid")
font_low = Font(name=FONT_FAMILY, size=9, color="166534")

fill_yes = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
font_yes = Font(name=FONT_FAMILY, size=9, bold=True, color="166534")

fill_supp = PatternFill(start_color="E0E7FF", end_color="E0E7FF", fill_type="solid")
font_supp = Font(name=FONT_FAMILY, size=9, color="3730A3")

fill_no = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
font_no = Font(name=FONT_FAMILY, size=9, color="64748B")

# Borders
border_thin = Border(left=Side(style="thin", color="CBD5E1"),
                     right=Side(style="thin", color="CBD5E1"),
                     top=Side(style="thin", color="CBD5E1"),
                     bottom=Side(style="thin", color="CBD5E1"))

border_double_bottom = Border(left=Side(style="thin", color="CBD5E1"),
                              right=Side(style="thin", color="CBD5E1"),
                              top=Side(style="thin", color="CBD5E1"),
                              bottom=Side(style="double", color="0F2537"))

border_card = Border(left=Side(style="medium", color="0F2537"),
                     right=Side(style="thin", color="CBD5E1"),
                     top=Side(style="thin", color="CBD5E1"),
                     bottom=Side(style="thin", color="CBD5E1"))

def auto_fit_columns(ws, min_width=12, max_width=60, custom_widths=None):
    if custom_widths is None:
        custom_widths = {}
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        if col_letter in custom_widths:
            ws.column_dimensions[col_letter].width = custom_widths[col_letter]
            continue
        max_len = 0
        for cell in col:
            # Skip title row and merged rows when auto-measuring
            if cell.row in [1, 2]:
                continue
            val_str = str(cell.value or "")
            if cell.hyperlink:
                val_str = "https://eur-lex.europa.eu"
            if len(val_str) > max_len:
                max_len = len(val_str)
        # Apply bounds
        w = max(min_width, min(max_len + 3, max_width))
        ws.column_dimensions[col_letter].width = w

print("Setup completed. Building sheets...")
